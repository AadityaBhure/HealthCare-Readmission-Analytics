"""Report engine: a tiny document model rendered to both DOCX (python-docx) and PDF (reportlab).

Content is assembled in reports/generate_report.py from pipeline artifacts; this module only
handles loading artifacts, numbering figures/tables and rendering.
Inline markup supported in paragraphs: **bold** and *italic*.
"""
from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from html import escape
from pathlib import Path

import pandas as pd

from . import config

# --------------------------------------------------------------------------------------
# Artifacts
# --------------------------------------------------------------------------------------


def load_artifacts(root: Path | None = None) -> dict:
    root = root or config.PROJECT_ROOT
    a, p = root / "artifacts", root / "data" / "processed"
    ids = {"hospital_id": str, "Hospital_ID": str}

    def csv(rel, base=a):
        return pd.read_csv(base / rel, dtype=ids)

    def js(rel):
        return json.loads((a / rel).read_text(encoding="utf-8"))

    art = {
        "audit": js("dataset_audit.json"), "findings": js("key_findings.json"), "manifest": js("pipeline_manifest.json"),
        "integrity": js("model_integrity.json"), "checks": csv("validation_checks.csv"), "prep_log": csv("preprocessing_log.csv"),
        "pop_summary": js("population/summary.json"), "pop_state": csv("population/statewide.csv"),
        "pop_county": csv("population/county.csv"), "pop_strata": csv("population/strata.csv"),
        "synthetic_summary": js("synthetic/summary.json"), "synthetic_predictions": csv("synthetic/scenario_predictions.csv"),
        "cabg_panel": csv("cabg_panel.csv", p), "stroke_panel": csv("stroke_panel.csv", p),
        "cabg_statewide": csv("cabg/statewide.csv"), "stroke_statewide": csv("stroke/statewide.csv"),
        "stroke_descriptive": js("stroke/descriptive.json"), "stroke_county": csv("stroke/county_summary.csv"),
        "cabg_feature_sets": csv("cabg/feature_set_comparison.csv"),
        "cabg_hospital_predictions": csv("cabg/hospital_predictions.csv"),
        "cabg_schema": json.loads((root / "models" / "cabg_feature_schema.json").read_text(encoding="utf-8")),
        "stroke_schema": json.loads((root / "models" / "stroke_feature_schema.json").read_text(encoding="utf-8")),
    }
    for pre in ("cabg", "stroke"):
        art[f"{pre}_summary"] = js(f"{pre}/summary.json")
        for key, rel in (("metrics", "metrics_test.csv"), ("cv_pooled", "cv_pooled_summary.csv"),
                         ("cv_nonoverlap", "cv_nonoverlap_summary.csv"), ("paired_cv", "paired_cv_vs_persistence.csv"),
                         ("bootstrap", "bootstrap_mae_diff.csv"), ("permutation", "permutation_importance.csv"),
                         ("native", "native_importance.csv"), ("error_by_volume", "error_by_volume.csv"),
                         ("learned_curve", "learned_curve.csv"), ("test_predictions", "test_predictions.csv"),
                         ("transitions", "transitions.csv"), ("selection", "train_cv_selection.csv"),
                         ("excluded", "excluded_transitions.csv")):
            art[f"{pre}_{key}"] = csv(f"{pre}/{rel}")
    from . import cabg_model  # local import to avoid a heavy import at module load
    art["cabg_transitions_nonoverlap"] = cabg_model.build_cabg_transitions(art["cabg_panel"].assign(
        hospital_id=art["cabg_panel"]["hospital_id"].astype(str)), horizon=2)
    return art


# --------------------------------------------------------------------------------------
# Document model
# --------------------------------------------------------------------------------------
@dataclass
class Report:
    title: str
    subtitle: str
    meta: list[tuple[str, str]]
    blocks: list[dict] = field(default_factory=list)
    _figs: dict = field(default_factory=dict)
    _tabs: dict = field(default_factory=dict)

    # numbering (assigned on first reference so numbers follow reading order)
    def fig(self, key: str) -> int:
        return self._figs.setdefault(key, len(self._figs) + 1)

    def tab(self, key: str) -> int:
        return self._tabs.setdefault(key, len(self._tabs) + 1)

    def h1(self, text):
        self.blocks.append({"type": "h1", "text": text})

    def h2(self, text):
        self.blocks.append({"type": "h2", "text": text})

    def p(self, text):
        self.blocks.append({"type": "p", "text": text})

    def bullets(self, items):
        self.blocks.append({"type": "bullets", "items": list(items)})

    def numbered(self, items):
        self.blocks.append({"type": "numbered", "items": list(items)})

    def note(self, text):
        self.blocks.append({"type": "note", "text": text})

    def code(self, text):
        self.blocks.append({"type": "code", "text": text})

    def pagebreak(self):
        self.blocks.append({"type": "pagebreak"})

    def table(self, key, columns, rows, caption, widths=None, font=8.0):
        self.blocks.append({"type": "table", "num": self.tab(key), "columns": columns,
                            "rows": [[("" if v is None else str(v)) for v in r] for r in rows],
                            "caption": caption, "widths": widths, "font": font})

    def figure(self, key, path, caption, width_in=6.3):
        if path is None or not Path(path).exists():
            return
        self.blocks.append({"type": "figure", "num": self.fig(key), "path": str(path), "caption": caption,
                            "width": width_in})


_MARK = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def _runs(text: str):
    for part in _MARK.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            yield part[2:-2], True, False
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            yield part[1:-1], False, True
        else:
            yield part, False, False


# --------------------------------------------------------------------------------------
# DOCX renderer
# --------------------------------------------------------------------------------------
def render_docx(rep: Report, path: Path) -> Path:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Inches, Pt, RGBColor

    NAVY = RGBColor(0x0B, 0x25, 0x45)
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.3))
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.2)

    st = doc.styles
    st["Normal"].font.name = "Calibri"
    st["Normal"].font.size = Pt(10.5)
    st["Normal"].paragraph_format.space_after = Pt(5)
    st["Normal"].paragraph_format.line_spacing = 1.15
    for name, size in (("Heading 1", 15), ("Heading 2", 12)):
        s = st[name]
        s.font.name, s.font.size, s.font.bold, s.font.color.rgb = "Calibri", Pt(size), True, NAVY
        s.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        s.paragraph_format.space_after = Pt(5)
        s.paragraph_format.keep_with_next = True
    st["Caption"].font.size = Pt(9)
    st["Caption"].font.italic = True
    st["Caption"].font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    def field(par, instr):
        r = par.add_run()
        b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
        t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = instr
        s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate")
        ph = OxmlElement("w:t"); ph.text = "Right-click and choose Update Field to build the table of contents." if "TOC" in instr else "1"
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
        r._r.append(b); r._r.append(t); r._r.append(s); r._r.append(ph); r._r.append(e)

    def add_runs(par, text, size=None, color=None):
        for txt, bold, ital in _runs(text):
            run = par.add_run(txt)
            run.bold, run.italic = bold or None, ital or None
            if size:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = color

    def shade(cell_or_par, hex_fill, is_par=False):
        el = cell_or_par._p.get_or_add_pPr() if is_par else cell_or_par._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
        el.append(shd)

    # ---- title page
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BTech Healthcare Analytics - Mini Project Report"); r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x0E, 0x7F, 0xA0); r.bold = True
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(rep.title); r.font.size = Pt(22); r.bold = True; r.font.color.rgb = NAVY
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(rep.subtitle); r.font.size = Pt(12); r.italic = True
    for _ in range(4):
        doc.add_paragraph()
    for k, v in rep.meta:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{k}: "); r.bold = True
        p.add_run(v)
    p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)

    # ---- table of contents
    toc_title = doc.add_paragraph()  # plain paragraph so the TOC does not list itself
    r = toc_title.add_run("Table of Contents"); r.bold = True; r.font.size = Pt(15); r.font.color.rgb = NAVY
    field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
    settings = doc.settings.element
    upd = OxmlElement("w:updateFields"); upd.set(qn("w:val"), "true"); settings.append(upd)
    p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)

    # ---- footer page numbers
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("Page "); field(fp, "PAGE")

    usable_w = 21.0 - 4.6  # cm
    for b in rep.blocks:
        t = b["type"]
        if t in ("h1", "h2"):
            doc.add_paragraph(b["text"], style="Heading 1" if t == "h1" else "Heading 2")
        elif t == "p":
            par = doc.add_paragraph(); add_runs(par, b["text"])
            par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif t in ("bullets", "numbered"):
            for item in b["items"]:
                par = doc.add_paragraph(style="List Bullet" if t == "bullets" else "List Number")
                add_runs(par, item)
        elif t == "note":
            par = doc.add_paragraph(); add_runs(par, b["text"], size=9.5)
            shade(par, "EEF4FB", is_par=True)
            par.paragraph_format.left_indent = Cm(0.3); par.paragraph_format.right_indent = Cm(0.3)
        elif t == "code":
            for line in b["text"].splitlines():
                par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(0)
                run = par.add_run(line or " "); run.font.name = "Consolas"; run.font.size = Pt(8.5)
                shade(par, "F3F4F6", is_par=True)
            doc.add_paragraph()
        elif t == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif t == "figure":
            par = doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.keep_with_next = True
            par.add_run().add_picture(b["path"], width=Inches(min(b["width"], usable_w / 2.54)))
            cap = doc.add_paragraph(style="Caption"); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cap.add_run(f"Figure {b['num']}. "); r.bold = True
            add_runs(cap, b["caption"])
        elif t == "table":
            cap = doc.add_paragraph(style="Caption"); cap.paragraph_format.keep_with_next = True
            r = cap.add_run(f"Table {b['num']}. "); r.bold = True
            add_runs(cap, b["caption"])
            cols = b["columns"]
            tbl = doc.add_table(rows=1, cols=len(cols)); tbl.style = "Table Grid"; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            widths = b["widths"] or [1] * len(cols)
            ws = [Cm(usable_w * w / sum(widths)) for w in widths]
            for i, c in enumerate(cols):
                cell = tbl.rows[0].cells[i]; cell.width = ws[i]
                cell.paragraphs[0].text = ""
                run = cell.paragraphs[0].add_run(c); run.bold = True; run.font.size = Pt(b["font"]); run.font.color.rgb = RGBColor(255, 255, 255)
                shade(cell, "0B2545")
            for ri, row in enumerate(b["rows"]):
                cells = tbl.add_row().cells
                for i, v in enumerate(row):
                    cells[i].width = ws[i]
                    par = cells[i].paragraphs[0]; par.paragraph_format.space_after = Pt(0)
                    add_runs(par, v, size=b["font"])
                    if ri % 2:
                        shade(cells[i], "F5F7FA")
            doc.add_paragraph()
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


# --------------------------------------------------------------------------------------
# PDF renderer
# --------------------------------------------------------------------------------------
def _register_fonts() -> str:
    import matplotlib
    from reportlab.lib.fonts import addMapping
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    ttf = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
    try:
        for name, file in (("DejaVu", "DejaVuSans.ttf"), ("DejaVu-Bold", "DejaVuSans-Bold.ttf"),
                           ("DejaVu-Oblique", "DejaVuSans-Oblique.ttf"), ("DejaVu-BoldOblique", "DejaVuSans-BoldOblique.ttf"),
                           ("DejaVuMono", "DejaVuSansMono.ttf")):
            pdfmetrics.registerFont(TTFont(name, str(ttf / file)))
        addMapping("DejaVu", 0, 0, "DejaVu"); addMapping("DejaVu", 1, 0, "DejaVu-Bold")
        addMapping("DejaVu", 0, 1, "DejaVu-Oblique"); addMapping("DejaVu", 1, 1, "DejaVu-BoldOblique")
        return "DejaVu"
    except Exception:  # pragma: no cover - fallback if matplotlib fonts are unavailable
        return "Helvetica"


def _rl_markup(text: str) -> str:
    out = []
    for txt, bold, ital in _runs(text):
        t = escape(txt)
        out.append(f"<b>{t}</b>" if bold else f"<i>{t}</i>" if ital else t)
    return "".join(out)


def render_pdf(rep: Report, path: Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm, inch
    from reportlab.platypus import (BaseDocTemplate, Frame, Image, KeepTogether, ListFlowable, ListItem, PageBreak,
                                    PageTemplate, Paragraph, Spacer, Table, TableStyle, XPreformatted)
    from reportlab.platypus.tableofcontents import TableOfContents

    font = _register_fonts()
    bold = "DejaVu-Bold" if font == "DejaVu" else "Helvetica-Bold"
    mono = "DejaVuMono" if font == "DejaVu" else "Courier"
    navy, teal, grey = colors.HexColor("#0B2545"), colors.HexColor("#0E7FA0"), colors.HexColor("#475569")
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName=font, fontSize=9.6, leading=13.6, alignment=TA_JUSTIFY, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=body, fontName=bold, fontSize=14.5, leading=18, textColor=navy, spaceBefore=12, spaceAfter=6, alignment=0)
    h2 = ParagraphStyle("h2", parent=h1, fontSize=11.5, leading=15, spaceBefore=9, spaceAfter=4)
    cap = ParagraphStyle("cap", parent=body, fontSize=8.4, leading=11, textColor=grey, alignment=TA_CENTER, spaceBefore=3, spaceAfter=10)
    tcap = ParagraphStyle("tcap", parent=cap, alignment=0, spaceBefore=6, spaceAfter=4)
    note = ParagraphStyle("note", parent=body, fontSize=9, leading=12.5, backColor=colors.HexColor("#EEF4FB"),
                          borderPadding=(6, 7, 6, 7), borderColor=colors.HexColor("#0E7FA0"), borderWidth=0, spaceBefore=6, spaceAfter=10)
    cell = ParagraphStyle("cell", parent=body, fontSize=7.6, leading=9.6, alignment=0, spaceAfter=0)
    head = ParagraphStyle("head", parent=cell, fontName=bold, textColor=colors.white)
    codes = ParagraphStyle("code", parent=body, fontName=mono, fontSize=7.8, leading=10, backColor=colors.HexColor("#F3F4F6"),
                           borderPadding=5, alignment=0)

    class Doc(BaseDocTemplate):
        def afterFlowable(self, flowable):
            if isinstance(flowable, Paragraph) and flowable.style.name in ("h1", "h2"):
                level = 0 if flowable.style.name == "h1" else 1
                text = flowable.getPlainText()
                key = f"h{id(flowable)}"
                self.canv.bookmarkPage(key)
                self.notify("TOCEntry", (level, text, self.page, key))

    def on_page(canv, doc):
        if doc.page == 1:
            return
        canv.saveState()
        canv.setFont(font, 7.5); canv.setFillColor(grey)
        canv.drawString(2.2 * cm, 1.2 * cm, rep.title[:95])
        canv.drawRightString(A4[0] - 2.2 * cm, 1.2 * cm, f"Page {doc.page}")
        canv.setStrokeColor(colors.HexColor("#E4E7EC")); canv.line(2.2 * cm, 1.5 * cm, A4[0] - 2.2 * cm, 1.5 * cm)
        canv.restoreState()

    doc = Doc(str(path), pagesize=A4, leftMargin=2.2 * cm, rightMargin=2.2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
              title=rep.title, author="Project team")
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="f")
    doc.addPageTemplates([PageTemplate(id="p", frames=[frame], onPage=on_page)])
    W = doc.width

    story = [Spacer(1, 4.5 * cm),
             Paragraph("BTech Healthcare Analytics - Mini Project Report",
                       ParagraphStyle("k", parent=body, fontName=bold, textColor=teal, alignment=TA_CENTER, fontSize=11)),
             Spacer(1, 0.4 * cm),
             Paragraph(escape(rep.title), ParagraphStyle("t", parent=h1, fontSize=20, leading=26, alignment=TA_CENTER)),
             Spacer(1, 0.3 * cm),
             Paragraph(f"<i>{escape(rep.subtitle)}</i>", ParagraphStyle("s", parent=body, alignment=TA_CENTER, fontSize=11)),
             Spacer(1, 3 * cm)]
    for k, v in rep.meta:
        story.append(Paragraph(f"<b>{escape(k)}:</b> {escape(v)}", ParagraphStyle("m", parent=body, alignment=TA_CENTER)))
    story.append(PageBreak())
    toc = TableOfContents()
    toc.levelStyles = [ParagraphStyle("toc1", fontName=bold, fontSize=9.5, leading=13, leftIndent=0, firstLineIndent=0),
                       ParagraphStyle("toc2", fontName=font, fontSize=8.8, leading=11.5, leftIndent=14, firstLineIndent=0)]
    story += [Paragraph("Table of Contents", ParagraphStyle("tch", parent=h1)), toc, PageBreak()]

    for b in rep.blocks:
        t = b["type"]
        if t == "h1":
            story.append(Paragraph(escape(b["text"]), h1))
        elif t == "h2":
            story.append(Paragraph(escape(b["text"]), h2))
        elif t == "p":
            story.append(Paragraph(_rl_markup(b["text"]), body))
        elif t in ("bullets", "numbered"):
            items = [ListItem(Paragraph(_rl_markup(i), ParagraphStyle("li", parent=body, spaceAfter=2)), leftIndent=12)
                     for i in b["items"]]
            story.append(ListFlowable(items, bulletType="bullet" if t == "bullets" else "1", start="•" if t == "bullets" else None,
                                      leftIndent=14, bulletFontName=font, bulletFontSize=8))
            story.append(Spacer(1, 4))
        elif t == "note":
            story.append(Paragraph(_rl_markup(b["text"]), note))
        elif t == "code":
            story.append(XPreformatted(escape(b["text"]), codes)); story.append(Spacer(1, 8))
        elif t == "pagebreak":
            story.append(PageBreak())
        elif t == "figure":
            img = Image(b["path"])
            w = min(b["width"] * inch, W)
            img.drawHeight = img.imageHeight * w / img.imageWidth
            img.drawWidth = w
            if img.drawHeight > 20 * cm:
                img.drawWidth *= 20 * cm / img.drawHeight; img.drawHeight = 20 * cm
            story.append(KeepTogether([img, Paragraph(f"<b>Figure {b['num']}.</b> " + _rl_markup(b["caption"]), cap)]))
        elif t == "table":
            widths = b["widths"] or [1] * len(b["columns"])
            cw = [W * x / sum(widths) for x in widths]
            cst = ParagraphStyle("c2", parent=cell, fontSize=b["font"] - 0.4, leading=(b["font"] - 0.4) * 1.27)
            hst = ParagraphStyle("h2c", parent=head, fontSize=b["font"] - 0.4, leading=(b["font"] - 0.4) * 1.27)
            data = [[Paragraph(escape(c), hst) for c in b["columns"]]]
            data += [[Paragraph(_rl_markup(v), cst) for v in r] for r in b["rows"]]
            tb = Table(data, colWidths=cw, repeatRows=1)
            style = [("BACKGROUND", (0, 0), (-1, 0), navy), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                     ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
                     ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
                     ("LEFTPADDING", (0, 0), (-1, -1), 3.5), ("RIGHTPADDING", (0, 0), (-1, -1), 3.5)]
            style += [("BACKGROUND", (0, i), (-1, i), colors.HexColor("#F5F7FA")) for i in range(2, len(data), 2)]
            tb.setStyle(TableStyle(style))
            story.append(Paragraph(f"<b>Table {b['num']}.</b> " + _rl_markup(b["caption"]), tcap))
            story.append(tb); story.append(Spacer(1, 8))
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.multiBuild(story)
    return path
